# Makes `openpyxl-rich.xlsx`: a workbook with everything a rewrite from rows would lose, for
# `xlsxFeatures`. Run from the package directory:
#   python3 -m venv /tmp/m365venv-x && /tmp/m365venv-x/bin/pip install openpyxl
#   /tmp/m365venv-x/bin/python __tests__/fixtures/make-openpyxl-rich.py
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "Budget"
ws.append(["Item", "Q1", "Q2", "Total"])
for i, (n, a, b) in enumerate([("Rent", 1200, 1250), ("Travel", 300.5, 410.25), ("Software", 99, 129)], start=2):
    ws.append([n, a, b, f"=SUM(B{i}:C{i})"])
ws["A1"].font = Font(bold=True, color="FF0000")
ws["B2"].number_format = "$#,##0.00"
ws.merge_cells("A6:D6"); ws["A6"] = "Notes"
ws["A7"].hyperlink = "https://example.com"; ws["A7"] = "link"
ws.add_table(Table(displayName="Costs", ref="A1:D4"))
ws.conditional_formatting.add("D2:D4", CellIsRule(operator="greaterThan", formula=["1000"], fill=PatternFill("solid", fgColor="FFC7CE")))
dv = DataValidation(type="list", formula1='"Yes,No"'); ws.add_data_validation(dv); dv.add("E2")
ws["A2"].comment = Comment("Paid monthly", "Phillip")
chart = BarChart(); chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True); ws.add_chart(chart, "F2")
wb.defined_names["Rent"] = DefinedName("Rent", attr_text="Budget!$B$2")
s2 = wb.create_sheet("Notes"); s2["A1"] = "Second sheet"
wb.save("__tests__/fixtures/openpyxl-rich.xlsx")
